from __future__ import annotations
import re
from .models import Clause

ARTICLE = re.compile(r'(?m)^\s*(제\s*\d+\s*조(?:\s*의\s*\d+)?(?:\([^\n)]{1,150}\))?|부\s*칙(?:\([^\n)]{1,100}\))?|별\s*표(?:\([^\n)]{1,100}\))?)\s*')
ARTICLE_ID = re.compile(r'(제\s*\d+\s*조(?:\s*의\s*\d+)?)')

def clean(text: str) -> str:
    return re.sub(r'\s+', ' ', text).strip()

def parse_clauses(text: str) -> list[Clause]:
    if not text or not text.strip():
        raise ValueError('문서에 읽을 수 있는 텍스트가 없습니다.')
    matches = list(ARTICLE.finditer(text))
    if not matches:
        return [Clause('전체문서', '전체 문서', text.strip())]
    out: list[Clause] = []
    prefix = text[:matches[0].start()].strip()
    if prefix:
        out.append(Clause('전문', '서문·기본정보', prefix))
    counts: dict[str,int] = {}
    in_attachment = False
    for idx, match in enumerate(matches):
        heading = clean(match.group(1))
        base_match = ARTICLE_ID.match(heading)
        # Internal 제1조 headings within a 부칙/별표 belong to that section.
        if in_attachment and base_match:
            continue
        if not base_match:
            in_attachment = True
        base = re.sub(r'\s+', '', base_match.group(1)) if base_match else heading.replace(' ', '')
        counts[base] = counts.get(base, 0) + 1
        identifier = f'{base}#{counts[base]}' if counts[base] > 1 else base
        next_index = idx + 1
        if not base_match:
            while next_index < len(matches) and ARTICLE_ID.match(matches[next_index].group(1)):
                next_index += 1
        end = matches[next_index].start() if next_index < len(matches) else len(text)
        body = text[match.end():end].strip()
        out.append(Clause(identifier, heading, body))
    return out


def parse_bytes(raw: bytes, filename: str, allow_ocr: bool=False) -> str:
    import io
    if len(raw) > 12 * 1024 * 1024:
        raise ValueError('입력 문서 크기 제한 12MB를 초과했습니다.')
    ext = filename.lower().rsplit('.', 1)[-1]
    if ext in ('txt', 'md'):
        for enc in ('utf-8-sig', 'cp949'):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                pass
        raise ValueError('텍스트 인코딩을 읽지 못했습니다. UTF-8로 저장해 주세요.')
    if ext == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if len(reader.pages) > 200:
            raise ValueError('PDF 페이지 제한(200)을 초과했습니다.')
        extracted = '\n'.join(p.extract_text() or '' for p in reader.pages)
        if not extracted.strip():
            if not allow_ocr:
                raise ValueError('스캔 PDF로 추정됩니다. OCR을 활성화하려면 선택 옵션과 로컬 Tesseract(한국어 언어팩)가 필요합니다.')
            try:
                import fitz, pytesseract
                doc=fitz.open(stream=raw,filetype='pdf')
                if len(doc)>40:
                    raise ValueError('OCR은 안전 및 성능상 40페이지까지만 지원합니다.')
                extracted='\n'.join(pytesseract.image_to_string(
                    __import__('PIL.Image',fromlist=['Image']).frombytes('RGB',[pix.width,pix.height],pix.samples),lang='kor+eng')
                    for page in doc for pix in [page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)])
            except ValueError:
                raise
            except Exception as e:
                raise ValueError('OCR 실행 실패: Tesseract 프로그램 및 kor 언어팩 설치 여부를 확인하세요.') from e
            if not extracted.strip():
                raise ValueError('OCR에서도 텍스트를 추출하지 못했습니다.')
        return extracted
    if ext == 'docx':
        from docx import Document as DocxDocument
        doc = DocxDocument(io.BytesIO(raw))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(' | '.join(c.text for c in row.cells))
        return '\n'.join(paragraphs)
    if ext == 'hwpx':
        import zipfile, re
        from defusedxml import ElementTree as ET
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                names=[n for n in archive.namelist() if re.fullmatch(r'Contents/section[0-9]+\.xml',n,re.I)]
                if not names: raise ValueError('HWPX 섹션 XML이 없습니다.')
                if len(names)>200: raise ValueError('HWPX 섹션이 너무 많습니다.')
                paragraphs=[]
                for name in sorted(names,key=lambda x:int(re.search(r'(\d+)\.xml',x).group(1))):
                    info=archive.getinfo(name)
                    if info.file_size>15*1024*1024: raise ValueError('HWPX 섹션 XML 제한 초과')
                    root=ET.fromstring(archive.read(name))
                    for p in root.iter():
                        if p.tag.endswith('}p') or p.tag=='p':
                            parts=[(t.text or '') for t in p.iter() if t.tag.endswith('}t') or t.tag=='t']
                            if parts: paragraphs.append(''.join(parts))
                if paragraphs: return '\n'.join(paragraphs)
                raise ValueError('HWPX 본문 텍스트를 찾지 못했습니다.')
        except (zipfile.BadZipFile,KeyError) as exc:
            raise ValueError('손상되었거나 지원되지 않는 HWPX 파일입니다.') from exc
    if ext == 'hwp':
        raise ValueError('구형 HWP 바이너리는 직접 읽지 않습니다. 한글에서 HWPX/DOCX/PDF로 변환하세요.')
    raise ValueError('지원하지 않는 형식입니다. TXT/MD/PDF/DOCX/HWPX를 사용해 주세요.')
